"""Exercise export correctness with synthetic records and no database access."""

import asyncio
import csv
import io
import json
from zipfile import ZipFile

import httpx
import pytest
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook
from docx import Document

from server import main


DETAILS = {
    "generated_at": "2026-10-28T01:00:00+00:00",
    "source": "Synchronized Supabase records: level_attempts",
    "sample_size": 1,
    "filters": {"date_from": "2026-10-01", "date_to": "2026-10-28"},
}


def filters(dataset="attempts", **values):
    return main._export_filters(dataset, values.get("date_from"), values.get("date_to"),
        values.get("week"), values.get("difficulty"), values.get("stratum"), values.get("include_incomplete", True))


def test_date_range_keeps_both_bounds_and_includes_fractional_last_second():
    selected, applied = filters(date_from="2026-10-01", date_to="2026-10-28", include_incomplete=False)
    assert selected == {
        "and": "(started_at.gte.2026-10-01T00:00:00Z,started_at.lt.2026-10-29T00:00:00Z)",
        "status": "not.in.(incomplete,abandoned,restarted)",
    }
    assert applied == {"date_from": "2026-10-01", "date_to": "2026-10-28", "include_incomplete": False}
    single, _ = filters(date_to="2026-12-31")
    assert single["and"] == "(started_at.lt.2027-01-01T00:00:00Z)"


@pytest.mark.parametrize("values", [
    {"date_from": "20261001"}, {"date_to": "2026-W43-1"}, {"date_to": "2026-02-30"},
    {"date_from": "2026-10-29", "date_to": "2026-10-28"}, {"date_to": "9999-12-31"},
    {"week": 10}, {"difficulty": "unknown"}, {"stratum": "Unknown"},
])
def test_invalid_filters_fail_before_fetch(values):
    with pytest.raises(HTTPException) as failure:
        filters(**values)
    assert failure.value.status_code == 422


def test_inapplicable_filters_are_rejected_instead_of_silently_ignored():
    for dataset, values in [("weekly", {"date_from": "2026-10-01"}), ("events", {"week": 1}),
                            ("coins", {"difficulty": "standard"}), ("attempts", {"stratum": "Minimal"})]:
        with pytest.raises(HTTPException) as failure:
            filters(dataset, **values)
        assert failure.value.status_code == 422
    selected, applied = filters("participants", week=2, stratum="Minimal")
    assert selected == {"current_week": "eq.2", "stratum": "eq.Minimal"}
    assert applied == {"week": 2, "stratum": "Minimal"}


@pytest.mark.parametrize("attack", ["=SUM(1,2)", " +cmd", "\t=cmd", "\n@SUM(A1)", "\u00a0-1+2", "\rplain"])
def test_spreadsheet_text_cannot_execute_as_a_formula(attack):
    safe = main._spreadsheet_value(attack)
    assert safe == "'" + attack
    assert main._spreadsheet_value(0) == 0
    assert main._spreadsheet_value(False) is False
    assert main._spreadsheet_value(None) == ""
    assert main._spreadsheet_value("") == ""
    assert main._spreadsheet_value(-6) == -6


def test_csv_preserves_missing_false_zero_structured_values_and_utf8():
    rows = [{"participant_id": "LC-TEST", "score": 0, "completed": False, "missing": None,
             "text": "é\nline", "attack": " =1+1", "details": {"hints": 0}}]
    body, media = main._render_export("csv", "attempts", rows, DETAILS)
    assert body.startswith(b"\xef\xbb\xbf")
    assert media == "text/csv; charset=utf-8"
    result = next(csv.DictReader(io.StringIO(body.decode("utf-8-sig"))))
    assert result["score"] == "0" and result["completed"] == "False" and result["missing"] == ""
    assert result["attack"] == "' =1+1" and result["text"] == "é\nline"
    assert json.loads(result["details"]) == {"hints": 0}


def test_xlsx_keeps_numeric_types_and_escapes_formulas_and_invalid_xml_controls():
    rows = [{"score": 0, "completed": False, "missing": None, "text": "bad\x00text", "attack": " =1+1"}]
    body, _ = main._render_export("xlsx", "attempts", rows, DETAILS)
    workbook = load_workbook(io.BytesIO(body), data_only=False)
    sheet = workbook["attempts"]
    result = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
    assert result == {"score": 0, "completed": False, "missing": None,
                      "text": "bad\\u0000text", "attack": "' =1+1"}
    assert all(cell.data_type != "f" for cell in sheet[2])
    assert workbook.sheetnames == ["attempts", "Report Details"]
    assert any(row[0].value == "Applied filters" for row in workbook["Report Details"])


def test_word_preserves_zero_false_and_report_metadata():
    body, _ = main._render_export("docx", "attempts", [{"score": 0, "completed": False, "missing": None}], DETAILS)
    document = Document(io.BytesIO(body))
    values = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "0" in values and "False" in values and "" in values
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert DETAILS["source"] in paragraphs and "2026-10-28" in paragraphs


@pytest.mark.parametrize("file_format", ["csv", "xlsx", "pdf", "docx"])
def test_empty_dataset_has_a_valid_export(file_format):
    body, media = main._render_export(file_format, "attempts", [], {**DETAILS, "sample_size": 0})
    assert body and media
    if file_format == "pdf":
        assert body.startswith(b"%PDF-") and body.rstrip().endswith(b"%%EOF")
    if file_format in {"xlsx", "docx"}:
        with ZipFile(io.BytesIO(body)) as archive:
            assert archive.testzip() is None


def test_export_paginates_and_replaces_internal_id_without_mutating_records(monkeypatch):
    raw = [{"student_id": "student-internal", "score": 0}, {"student_id": None, "score": 1}]
    calls = []
    monkeypatch.setenv("LOGICADE_PSEUDONYM_KEY", "test-only-pseudonym-key")

    async def query(_base, _key, _resource, **kwargs):
        calls.append(kwargs)
        if kwargs["offset"] == 0:
            return raw[:1], 2
        return raw[1:], 2

    monkeypatch.setattr(main, "_admin_rows", query)
    result = asyncio.run(main._export_rows("unused", "unused", "level_attempts", "started_at.asc", {"week": "eq.1"}))
    assert [call["offset"] for call in calls] == [0, 1]
    assert all(call["params"] == {"order": "started_at.asc", "week": "eq.1"} for call in calls)
    assert result[0]["participant_id"] == main._research_pseudonym("student-internal")
    assert result[1]["participant_id"] is None
    assert all("student_id" not in row for row in result)
    assert raw[0]["student_id"] == "student-internal" and "participant_id" not in raw[0]


def test_route_records_actual_filters_source_actor_and_sample_before_download(monkeypatch):
    fetched, audit = [], []
    monkeypatch.setattr(main, "_check_admin_key", lambda _: ("unused", "unused"))
    monkeypatch.setattr(main, "_admin_audit_actor", lambda token: "admin@example.test" if token == "Bearer admin-session" else "unexpected")

    async def rows(_base, _key, resource, order, selected):
        fetched.append((resource, selected))
        return [{"participant_id": "LC-1", "score": 0}, {"participant_id": "LC-1", "score": 1}]

    async def record(*args, **kwargs):
        audit.append((args, kwargs))

    monkeypatch.setattr(main, "_export_rows", rows)
    monkeypatch.setattr(main, "_record_admin_audit", record)
    app = FastAPI()
    app.get("/report.{file_format}")(main.admin_export_document)
    app.get("/legacy.csv")(main.admin_export_csv)

    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get("/report.csv?dataset=attempts&date_from=2026-10-01&date_to=2026-10-28&include_incomplete=false",
                headers={"X-Admin-Key": "Bearer admin-session"})
            assert response.status_code == 200
            assert response.headers["cache-control"] == "no-store"
            legacy = await client.get("/legacy.csv?dataset=participants", headers={"X-Admin-Key": "Bearer admin-session"})
            assert legacy.status_code == 200
    asyncio.run(run())
    assert fetched[0] == ("level_attempts", filters(date_from="2026-10-01", date_to="2026-10-28", include_incomplete=False)[0])
    args, kwargs = audit[0]
    assert args[2:4] == ("report.export.csv", "attempts")
    assert args[4]["filters"] == {"date_from": "2026-10-01", "date_to": "2026-10-28", "include_incomplete": False}
    assert args[4]["source"] == "level_attempts" and args[4]["sample_size"] == 1 and args[4]["row_count"] == 2
    assert kwargs == {"actor_id": "admin@example.test", "required": True}
    assert fetched[1] == ("v_admin_participants", {})


@pytest.mark.parametrize("failure", ["response", "network"])
def test_required_audit_failure_prevents_an_unlogged_export(monkeypatch, failure):
    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            pass

        async def post(self, url, **_kwargs):
            if failure == "network":
                raise httpx.ConnectError("test-only unavailable")
            return httpx.Response(503, text="test-only unavailable", request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "AsyncClient", Client)
    with pytest.raises(HTTPException) as blocked:
        asyncio.run(main._record_admin_audit("http://unused", "unused", "report.export.csv", "attempts", {}, required=True))
    assert blocked.value.status_code == 503
    # Existing best-effort callers, such as account login, retain their policy.
    asyncio.run(main._record_admin_audit("http://unused", "unused", "admin.login", "admin", {}))
